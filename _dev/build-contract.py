"""Build the editable Czech contract draft; PDF is exported from this DOCX.

Run with the bundled Python runtime (python-docx). No external data or signing.
"""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'assets' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
# Remove Word template border residue, including the default Title rule.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.68), Inches(.65)
section.left_margin = section.right_margin = Inches(.78)
section.header_distance = section.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().append(OxmlElement('w:lang'))
    style.element.rPr.find(qn('w:lang')).set(qn('w:val'), 'cs-CZ')
    style.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(21)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(9)
doc.styles['Subtitle'].font.size = Pt(10)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
for name, size in [('Heading 1', 13), ('Heading 2', 11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(13)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = 'Smlouva o zajištění hudební produkce BoneSaver'
doc.core_properties.subject = 'Pracovní vzor smlouvy pro jednotlivé hudební akce'
doc.core_properties.author = 'BoneSaver'
doc.core_properties.last_modified_by = 'BoneSaver'
doc.core_properties.created = datetime(2026, 9, 24, tzinfo=timezone.utc)
doc.core_properties.modified = datetime(2026, 9, 24, tzinfo=timezone.utc)
doc.core_properties.keywords = 'BoneSaver, smlouva, hudební produkce, vzor'

header = section.header.paragraphs[0]
header.text = 'BONESAVER    |    PRACOVNÍ VZOR SMLOUVY'
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Verze 24. 9. 2026    |    Strana ').font.size = Pt(8)
for instr in ['PAGE', 'NUMPAGES']:
    if instr == 'NUMPAGES': footer.add_run(' z ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instr)
    footer._p.append(field)

def p(text, bold=False, size=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    if size: run.font.size = Pt(size)
    return para

def heading(text):
    return doc.add_paragraph(text, 'Heading 1')

def newpage():
    doc.add_page_break()

def table(headers, rows, widths=(2.25, 4.69)):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: el.set(qn('w:' + key), val)
        borders.append(el)
    pr.append(borders)
    for row_index, texts in enumerate([headers] + rows):
        row = t.rows[0] if row_index == 0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if row_index == 0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, text, width in zip(row.cells, texts, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '90' if edge in ['top', 'bottom'] else '110')
                el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            if row_index == 0:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E7EDF3'); tcpr.append(shading)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(text); run.font.size = Pt(10.5); run.bold = row_index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc.add_paragraph('Smlouva o zajištění\nhudební produkce', 'Title')
doc.add_paragraph('BoneSaver    |    Vzor k individuálnímu doplnění', 'Subtitle')
p('Před podpisem doplňte všechna pole v hranatých závorkách, potvrďte obchodní podmínky a technickou přílohu. Tento nevyplněný vzor sám o sobě nerezervuje termín. Před prvním použitím doporučujeme právní kontrolu.', size=9.5)
p('Smluvní strany uzavírají podle § 1746 odst. 2 zákona č. 89/2012 Sb., občanského zákoníku, tuto smlouvu.', size=10)
heading('1 Smluvní strany')
p('Poskytovatel', bold=True)
p('Ing. Miroslav Bečička, IČO 69126887\nSídlo Do Polí 1054, 530 06 Pardubice\nDIČ CZ7512233322, plátce DPH\nE-mail bonesavermusic@gmail.com, telefon +420 734 393 711', size=10.5)
p('Poskytovatel zajišťuje vystoupení kapely BoneSaver a odpovídá objednateli za jeho sjednané provedení. Smluvní stranou není samotný název kapely.', size=10)
p('Objednatel', bold=True)
p('Jméno a příjmení / název: [DOPLNIT]\nAdresa bydliště / sídla: [DOPLNIT]\nIČO a DIČ, je-li relevantní: [DOPLNIT nebo NEPOUŽIJE SE]\nOsoba oprávněná jednat: [DOPLNIT nebo NEPOUŽIJE SE]\nE-mail a telefon: [DOPLNIT]', size=10.5)
heading('2 Akce a časový rozsah')
table(['Údaj', 'Dohodnutá hodnota'], [
    ['Název a typ akce', '[DOPLNIT]'],
    ['Místo a přesná adresa', '[DOPLNIT]'],
    ['Datum a čas produkce', '[DOPLNIT začátek i konec včetně data po půlnoci]'],
    ['Rozsah živého hraní', '[DOPLNIT počet bloků a jejich délku v minutách]'],
    ['Přestávky', '[DOPLNIT počet a délku, případně časový plán]'],
], widths=(2.15, 4.79))

newpage()
heading('3 Předmět a průběh produkce')
p('3.1 Poskytovatel zajistí živé hudební vystoupení kapely BoneSaver v rozsahu článku 2. Objednatel poskytne sjednanou součinnost a zaplatí cenu podle článku 4. Technické a organizační podmínky stanoví příloha 1, která je součástí smlouvy.')
p('3.2 Program vychází z repertoáru odsouhlaseného při uzavření smlouvy. Zvláštní požadavky, první tanec, vlastní tvorbu nebo nově nastudovanou píseň uveďte do přílohy. Píseň mimo potvrzený program není automaticky součástí závazku. Pořadí skladeb může kapela přizpůsobit průběhu akce.')
p('3.3 Prodloužení hraní, další technika a jiné služby nad sjednaný rozsah vyžadují dohodu obou stran o rozsahu i ceně před jejich poskytnutím. Zpoždění programu samo o sobě neposouvá sjednaný konec. Případné krácení, náhradu času a vliv na cenu strany zaznamenají podle skutečné příčiny.')
heading('4 Cena a platební podmínky')
table(['Položka', 'Částka nebo podmínka'], [
    ['Cena bez DPH', '[DOPLNIT] Kč'],
    ['DPH', 'Sazba [DOPLNIT] %, částka [DOPLNIT] Kč'],
    ['Celková cena včetně DPH', '[DOPLNIT] Kč'],
    ['Zahrnuto v celkové ceně', '[DOPLNIT dopravu, techniku a další dohodnuté služby]'],
    ['Případné další náklady', '[DOPLNIT přesně nebo ŽÁDNÉ]'],
    ['Záloha a její splatnost', '50 % celkové ceny, tedy [DOPLNIT] Kč\nDo 14 dnů od podpisu smlouvy oběma stranami'],
    ['Doplatek a splatnost', '50 % celkové ceny, tedy [DOPLNIT] Kč\nPo skončení sjednaného vystoupení'],
    ['Úhrada a účet', '[DOPLNIT převodem / hotově, účet a platební identifikaci]'],
])
p('4.1 Celková cena je sjednána včetně všech výslovně uvedených nákladů a DPH. Poskytovatel ji nemůže jednostranně zvýšit. Případnou změnu ceny je třeba předem výslovně odsouhlasit; u spotřebitele vždy s uvedením konečné částky.')
p('4.2 Objednatel uhradí zálohu ve výši 50 % celkové ceny včetně DPH do 14 dnů ode dne, kdy smlouvu podepíše poslední ze smluvních stran. Zbývající doplatek 50 % je splatný po skončení sjednaného vystoupení. Záloha je částí ceny, nikoli závdavkem; na storno se započte pouze v případě podle článku 6. Poskytovatel vystaví odpovídající doklad. Při převodu je rozhodující připsání částky na účet; při úhradě hotově poskytovatel předá potvrzení.')
p('4.3 Neuhrazení zálohy samo o sobě neruší smlouvu. Poskytovatel nejprve vyzve k úhradě a poskytne přiměřenou dodatečnou lhůtu. Další postup se řídí smlouvou a zákonem.')

newpage()
heading('5 Součinnost a bezpečnost')
p('5.1 Objednatel zajistí včasný přístup do místa konání, bezpečný prostor a elektrické napájení podle přílohy 1. Poskytovatel odpovídá za bezpečný stav a obsluhu vlastní techniky. Pořadatel předem oznámí limity hlučnosti, konec povolené produkce a další pravidla místa konání.')
p('5.2 Při bezprostředním ohrožení zdraví nebo techniky, zejména vodou, bouřkou, nebezpečným napájením nebo agresivním chováním, může kapela hraní v nezbytném rozsahu přerušit. Poskytovatel neprodleně oznámí důvod a umožní nápravu, je-li bezpečně možná. Obě strany předcházejí škodám. Přerušení nezakládá automatický nárok na plnou cenu; vypořádání závisí na příčině a skutečně poskytnutém plnění.')
heading('6 Zrušení akce a změna termínu')
p('6.1 Změna termínu vyžaduje souhlas obou stran. Objednatel může objednané vystoupení zrušit oznámením doručeným poskytovateli na smluvní e-mail. Při takovém dobrovolném zrušení ze strany objednatele činí sjednané storno 50 % celkové ceny včetně DPH podle článku 4, bez ohledu na dobu zbývající do akce.')
p('6.2 Zaplacená záloha se na storno započte; nehradí se tedy záloha a dalších 50 % navíc. Není-li záloha uhrazena nebo nepokrývá-li storno, objednatel doplatí rozdíl do 14 dnů od doručení oznámení o zrušení. Případný přeplatek poskytovatel vrátí ve stejné lhůtě. Za tutéž újmu se vedle sjednaného storna neúčtuje další paušál ani duplicitní náhrada.')
p('6.3 Storno se nepoužije při oprávněném odstoupení objednatele pro porušení smlouvy poskytovatelem ani tam, kde mu zákon dává právo závazek ukončit bez takové platby. U spotřebitele se ujednání použije pouze v rozsahu přípustném podle kogentních pravidel jeho ochrany.')
p('6.4 Poskytovatel může své vystoupení zrušit pouze při prokazatelných zdravotních problémech některého člena kapely, jestliže kvůli nim nelze vystoupení zajistit ani bez tohoto člena nebo s vhodnou náhradou. Musí nejprve vyvinout přiměřené úsilí k nalezení řešení, například vhodného záskoku, doporučení náhradní kapely nebo vystoupení v menším obsazení.')
p('6.5 Poskytovatel oznámí situaci bezodkladně. Zdravotní překážku doloží v nezbytném rozsahu, například potvrzením o nemožnosti vystoupení, bez požadavku na sdělení diagnózy. Jiná kapela nebo podstatně změněné obsazení vyžadují výslovný souhlas objednatele a dohodu o případné změně ceny. Doporučení jiné kapely samo nenahrazuje splnění smlouvy.')
p('6.6 Nedohodnou-li se strany na řešení a vystoupení se neuskuteční, poskytovatel vrátí veškeré přijaté úhrady za neposkytnuté plnění do 14 dnů od oznámení zrušení. Poskytovateli nevzniká právo na storno. Případné zákonné nároky objednatele zůstávají zachovány; zdravotní důvod neznamená automatické zproštění odpovědnosti za škodu.')
heading('7 Mimořádné překážky')
p('Zákonný zánik závazku pro nemožnost plnění a případné zproštění odpovědnosti za škodu se posuzují samostatně podle zákona; nejde o další smluvní právo kapely libovolně zrušit akci. Dotčená strana překážku bezodkladně oznámí, doloží a hledá řešení. Platby za neposkytnuté plnění se nevypořádávají automatickým propadnutím zálohy. Nízká účast ani běžně předvídatelné počasí nejsou automaticky vyšší mocí.', size=10)

newpage()
heading('8 Autorská práva a pořizování záznamů')
p('8.1 Objednatel jako pořadatel zajistí oprávnění k veřejnému užití hudebních děl a dalších chráněných předmětů, pokud je konkrétní akce vyžaduje, včetně příslušných licencí a odměn kolektivním správcům. Poskytovatel včas dodá potřebný seznam skladeb. Soukromý či veřejný režim akce a zahrnutí reprodukované hudby se upřesní v příloze; případné jiné rozdělení těchto povinností musí být výslovné.')
p('8.2 Smlouva sama neposkytuje licenci k veřejnému šíření, živému přenosu nebo komerčnímu využití záznamu výkonu kapely ani souhlas s použitím podoby hostů. Potřebná oprávnění a případné zákonné výjimky se posuzují samostatně. Ani kapela nezískává automatické oprávnění zveřejňovat fotografie hostů pro propagaci.')
heading('9 Vady plnění a odpovědnost')
p('9.1 Zjištěnou vadu je vhodné oznámit kontaktní osobě poskytovatele ihned, aby ji bylo možné napravit. Neuplatnění vady během akce samo o sobě neznamená vzdání se práv. Reklamaci lze uplatnit také na e-mailu nebo adrese poskytovatele uvedených v článku 1.')
p('9.2 Odpovědnost za škodu, vady plnění a nároky na slevu nebo vrácení ceny se řídí zákonem. Smlouva nevylučuje odpovědnost za úmysl, hrubou nedbalost nebo újmu na přirozených právech člověka. Objednatel neodpovídá automaticky za veškeré jednání hostů bez zákonného či sjednaného důvodu.')
heading('10 Ujednání pro spotřebitele')
p('Je-li objednatel spotřebitelem a jde o službu využití volného času na konkrétní datum či období podle článku 2, uplatní se výjimka z práva odstoupit bez důvodu podle § 1837 písm. j) občanského zákoníku. Tím nejsou dotčena práva z vad, porušení smlouvy ani sjednaná možnost storna. Pokud konkrétní plnění pod tuto výjimku nespadá, poskytovatel dodá příslušné zákonné poučení a formulář pro odstoupení před uzavřením smlouvy.')
p('Spotřebitelskou reklamaci poskytovatel vyřídí v zákonné lhůtě, zpravidla nejpozději do 30 dnů, pokud se se spotřebitelem nedohodne na delší lhůtě. K mimosoudnímu řešení spotřebitelského sporu je příslušná Česká obchodní inspekce, Štěpánská 567/15, 120 00 Praha 2, https://coi.gov.cz/informace-o-adr/.', size=10)
heading('11 Uzavření smlouvy a podpisy')
p('Smlouva se řídí českým právem a je uzavřena podpisem obou stran včetně přílohy 1. Změny smlouvy vyžadují písemný souhlas obou stran; e-mailová dohoda musí zachytit obsah změny a umožnit určit jednající osoby. Nezávazná poptávka nebo stažení vzoru nejsou přijetím smlouvy. Kontaktní údaje se používají k plnění smlouvy a zákonným povinnostem, nikoli automaticky k marketingu.', size=10)
p('Místo a datum podpisu: [DOPLNIT]', size=10)
p('Za poskytovatele: ____________________    Za objednatele: ____________________', size=10)
p('Ing. Miroslav Bečička                                  Jméno a funkce: [DOPLNIT]', size=10)

newpage()
doc.add_paragraph('Příloha 1\nTechnické a organizační podmínky', 'Title')
p('Nedílná součást smlouvy pro akci [NÁZEV] dne [DATUM]. Údaje níže obě strany potvrdí před podpisem. Pole, která se nepoužijí, označte NEPOUŽIJE SE.', size=10)
table(['Oblast', 'Dohodnuté podmínky'], [
    ['Kontakt na místě', 'Objednatel [jméno a telefon]\nPoskytovatel [jméno a telefon]'],
    ['Příjezd a příprava', 'Příjezd [čas], přístup do sálu [čas]\nZvuková zkouška [čas], odvoz techniky [čas]'],
    ['Prostor pro kapelu', 'Rozměry [šířka × hloubka], povrch [popis]\nZastřešení a ochrana proti vodě [popis]'],
    ['Elektrické napájení', '[Zásuvky, okruhy, jištění a vzdálenosti potvrzené poskytovatelem]'],
    ['Zvuk a světla', '[Kdo zajistí jakou techniku, obsluhu a rozsah ozvučení]'],
    ['Přístup a zázemí', '[Vykládka, schody, parkování, převlečení, voda a další dohoda]'],
    ['Režim akce a publikum', '[Soukromá / veřejná akce, uvnitř / venku, počet osob]'],
    ['Přestávky a mikrofon', '[Reprodukovaná hudba, proslovy, obsluha mikrofonu, licence]'],
    ['Limity a venkovní varianta', '[Hlukový limit, konec produkce, bezpečné náhradní místo]'],
    ['Zvláštní hudební požadavky', '[První tanec, písně, vlastní tvorba, nejzazší termín domluvy]'],
    ['Licence a záznamy', '[Licenční povinnosti pořadatele, případná dohoda o záznamu]'],
    ['Další dohodnuté podmínky', '[DOPLNIT nebo ŽÁDNÉ]'],
], widths=(2.0, 4.94))
p('Změny této přílohy musí být odsouhlaseny oběma stranami. Obecná informace na webu nenahrazuje konkrétní technické ujednání pro tuto akci.', size=10)
p('Za poskytovatele: ____________________    Za objednatele: ____________________', size=10)
p('Datum: [DOPLNIT]                                           Datum: [DOPLNIT]', size=10)

target = OUT / 'bonesaver-smlouva-hudebni-produkce-vzor.docx'
doc.save(target)
print(target)
